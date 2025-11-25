
# ============================================
# GENERAR EXCEL CON ESTADÍSTICAS DETALLADAS
# ============================================

# 1. Importar librerías necesarias
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from scipy.stats import f

# 2. Leer el archivo CSV
df = pd.read_csv("INDICADORES_MUNICIPALES.csv")

# 3. Filtrar columnas relevantes
df = df[["Tipo_Municipio", "Ingreso_Promedio_Mensual (MXN)", "Promedio_Escolaridad (años)"]]

# ============================================
# HOJA 1: ESTADÍSTICAS DETALLADAS
# ============================================

# 4. Calcular métricas básicas por tipo de municipio
stats = df.groupby("Tipo_Municipio")["Ingreso_Promedio_Mensual (MXN)"].agg([
    "count", "mean", "var", "std", "min", "max"
])

# 5. Calcular métricas adicionales
stats["rango"] = stats["max"] - stats["min"]
stats["coef_variacion"] = (stats["std"] / stats["mean"]) * 100

# 6. Calcular percentiles (25, 50, 75)
percentiles = df.groupby("Tipo_Municipio")["Ingreso_Promedio_Mensual (MXN)"].quantile([0.25, 0.5, 0.75]).unstack()
percentiles.columns = ["p25", "p50", "p75"]

# 7. Unir percentiles con stats
stats = stats.join(percentiles)

# 8. Calcular correlación escolaridad-ingreso por tipo de municipio
correlaciones = df.groupby("Tipo_Municipio").apply(
    lambda g: g["Ingreso_Promedio_Mensual (MXN)"].corr(g["Promedio_Escolaridad (años)"])
)
correlaciones = correlaciones.rename("correlacion_escolaridad_ingreso")

# ============================================
# CREAR ARCHIVO EXCEL
# ============================================

wb = Workbook()

# Hoja 1: Estadísticas Detalladas
ws1 = wb.active
ws1.title = "Estadísticas Detalladas"

# Encabezados
headers1 = [
    "Tipo de Municipio", "n", "Media", "Varianza", "Desv.Estándar",
    "Mínimo", "Máximo", "Rango", "Coef.Variación (%)",
    "p25", "p50", "p75", "Correlación Escolaridad-Ingreso"
]
ws1.append(headers1)

# Agregar datos
for tipo in stats.index:
    ws1.append([
        tipo,
        int(stats.loc[tipo, "count"]),
        round(stats.loc[tipo, "mean"], 2),
        round(stats.loc[tipo, "var"], 2),
        round(stats.loc[tipo, "std"], 2),
        round(stats.loc[tipo, "min"], 2),
        round(stats.loc[tipo, "max"], 2),
        round(stats.loc[tipo, "rango"], 2),
        round(stats.loc[tipo, "coef_variacion"], 2),
        round(stats.loc[tipo, "p25"], 2),
        round(stats.loc[tipo, "p50"], 2),
        round(stats.loc[tipo, "p75"], 2),
        round(correlaciones.loc[tipo], 4)
    ])

# Crear tabla en hoja 1
table1 = Table(displayName="TablaEstadisticasDetalladas", ref=f"A1:M{len(stats)+1}")
style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
table1.tableStyleInfo = style
ws1.add_table(table1)

# Ajustar ancho de columnas
for col in ws1.columns:
    max_length = max(len(str(cell.value)) for cell in col)
    ws1.column_dimensions[col[0].column_letter].width = max_length + 2

# ============================================
# HOJA 2: PRUEBA F
# ============================================

ws2 = wb.create_sheet(title="Prueba F")
headers2 = ["Comparación", "F", "gl", "p-value"]
ws2.append(headers2)

municipios = stats.index.tolist()
for i in range(len(municipios)):
    for j in range(i+1, len(municipios)):
        tipo1, tipo2 = municipios[i], municipios[j]
        var1, var2 = stats.loc[tipo1, "var"], stats.loc[tipo2, "var"]
        n1, n2 = stats.loc[tipo1, "count"], stats.loc[tipo2, "count"]

        # Calcular F y grados de libertad
        if var1 >= var2:
            F = var1 / var2
            df1, df2 = n1 - 1, n2 - 1
        else:
            F = var2 / var1
            df1, df2 = n2 - 1, n1 - 1

        # Calcular p-value
        p_value = 2 * min(f.cdf(F, df1, df2), 1 - f.cdf(F, df1, df2))

        ws2.append([f"{tipo1} vs {tipo2}", round(F, 3), f"({df1},{df2})", round(p_value, 4)])

# Crear tabla en hoja 2
table2 = Table(displayName="TablaPruebaF", ref=f"A1:D{len(ws2['A'])}")
table2.tableStyleInfo = style
ws2.add_table(table2)

# Ajustar ancho de columnas en hoja 2
for col in ws2.columns:
    max_length = max(len(str(cell.value)) for cell in col)
    ws2.column_dimensions[col[0].column_letter].width = max_length + 2

# Guardar archivo Excel
excel_file = "Resultados_Detallados.xlsx"
wb.save(excel_file)

print(f"Archivo Excel generado: {excel_file} con detalle extremo en dos hojas.")
